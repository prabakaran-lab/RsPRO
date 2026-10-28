import openpyxl
import os

files = [
    r"c:\Users\Kalai\Downloads\XSLX TO PDF\XSLX\Non_Contact_Safety_Switches_UPDATED (1) (1).xlsx",
    r"c:\Users\Kalai\Downloads\XSLX TO PDF\XSLX\Speciality Size Rechargeable Batteries (3) (1).xlsx"
]

for p in files:
    if not os.path.exists(p):
        print(f"File not found: {p}")
        continue
    wb = openpyxl.load_workbook(p, data_only=True)
    sheet = wb.active
    print(f"\nFile: {os.path.basename(p)}")
    print(f"Max row: {sheet.max_row}, Max col: {sheet.max_column}")
    row1 = [sheet.cell(1, col).value for col in range(1, 10)]
    row2 = [sheet.cell(2, col).value for col in range(1, 10)]
    print("  Row 1 (Roles):", row1)
    print("  Row 2 (Labels):", row2)
