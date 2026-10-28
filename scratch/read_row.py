import openpyxl

wb = openpyxl.load_workbook(r"c:\Users\Kalai\Downloads\XSLX TO PDF\XSLX\Non_Contact_Safety_Switches_UPDATED (1) (1).xlsx")
ws = wb.active

print("Headers (Row 1):")
headers = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column+1)]
print(headers)

print("\nData (Row 2):")
data = [ws.cell(row=2, column=c).value for c in range(1, ws.max_column+1)]
for i, (h, d) in enumerate(zip(headers, data)):
    h_clean = str(h).replace('\n', ' ')
    print(f"Col {i+1} [{h_clean}]: {d}")
